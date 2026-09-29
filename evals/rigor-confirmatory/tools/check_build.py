#!/usr/bin/env python3
"""Deterministic checks on built brains (v0.2 acceptance 3, 4a, 5; the confirmatory gate's code part).

For each business in the spec:
- the _brain tab exists and holds at least one owner note (source told);
- the saved copy holds exactly the source workbook's sheets and values plus _brain;
- question counts: substantive (the tool's own counter), all prompts;
- words in the owner's mouth: every content word of every told note comes from
  the question it answers (prompt, header, the labels and descriptions shown,
  the recommended label) or the owner's typed text, or the fixed note frame.

    .venv/bin/python check_build.py <spec.json> <src_root>
    spec.json: {"<business>": {"brain_files": [...], "home": "<SPREADSHEET_BRAIN_HOME>"}}
    src_root/<business>/<file name> is the untouched source workbook.
"""
import glob
import json
import os
import sys

REPO = os.path.expanduser("~/lab/spreadsheet-brain")
sys.path.insert(0, os.path.join(REPO, "skills", "spreadsheet-brain", "scripts"))
from sheetbrain import brainzip, interview  # noqa: E402
import openpyxl  # noqa: E402

# words the note frame itself adds around the owner's picks and words (brain._answer_notes, the build note)
# plus the tool's fixed consequence and coverage frames: "so these rows are left out of every count and total",
# "one of the owner's picks was", "so these values count as", "the owner said what <columns> mean"
FRAME = ("asked the owner picked wrote accepted recommended this question build first not sure was said and on "
         "also so these rows values left out of every count and total totals only one picks per what mean kept for who")


def answers(home):
    out = {}
    for st in sorted(glob.glob(os.path.join(home, "work", "*", "state.json"))):
        s = json.load(open(st))
        for qid, a in (s.get("answers") or {}).items():
            if isinstance(a, dict):
                out.setdefault(os.path.basename(os.path.dirname(st)), {})[qid] = a
        if isinstance(s.get("build"), dict):          # the build pick is kept beside the answers
            out.setdefault(os.path.basename(os.path.dirname(st)), {})["_build"] = s["build"]
    return out


def sheets(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        return {ws.title: [tuple(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}, wb.sheetnames
    finally:
        wb.close()


def same_data(src, out):
    s, s_names = sheets(src)
    o, o_names = sheets(out)
    extra = [n for n in o_names if n not in s_names]
    if extra != ["_brain"] or [n for n in o_names if n != "_brain"] != s_names:
        return False, f"sheets {s_names} -> {o_names}"
    for n in s_names:
        a, b = s[n], o[n]
        while a and not any(v not in (None, "") for v in a[-1]):
            a = a[:-1]
        while b and not any(v not in (None, "") for v in b[-1]):
            b = b[:-1]
        if a != b:
            bad = next((i for i, (x, y) in enumerate(zip(a, b)) if x != y), min(len(a), len(b)))
            return False, f"sheet {n!r} differs at row {bad + 1}"
    return True, ""


def lint(recs, by_q):
    flagged = []
    for r in recs:
        # owner notes only: edges and links the owner confirmed are code's sentences, not the owner's words
        if r.get("source") != "told" or r.get("record") not in ("fact", "open"):
            continue
        ref = str(r.get("ref") or "")
        qid = ref[2:] if ref.startswith("q:") else ""
        a = by_q.get(qid)
        if a is None and ref == "interview:coverage":
            # the coverage line names the columns the owner's answers were about
            a = {"prompt": " ".join(str(x.get("prompt") or "") + " " + str(x.get("header") or "") for x in by_q.values())}
        if a is None:
            flagged.append({"id": r.get("id"), "statement": r.get("statement"), "why": f"no answer for {ref!r}"})
            continue
        shown = " ".join([str(a.get("prompt") or ""), str(a.get("header") or "")])
        labels = " ".join([*(str(x) for x in a.get("labels") or []), str(a.get("recommended_label") or "")])
        descs = " ".join(str(v) for v in (a.get("descs") or {}).values())
        typed = str(a.get("text") or "")
        # a date the owner typed in full ('March 2') may be written short in the note ('Mar 2')
        months = "january february march april may june july august september october november december".split()
        # (the tool carries the owner's typed switch date into the notes of that switch's follow-ups,
        # so a month the owner typed in any answer counts)
        anywhere = " ".join(str(x.get("text") or "") + " " + str(x.get("prompt") or "") for x in by_q.values())
        typed += " " + " ".join(m[:3] for m in months if m in (typed + " " + anywhere).lower())
        unseen = interview.unseen_words(r.get("statement", ""), shown, labels, " ".join([descs, typed, FRAME]))
        if unseen:
            flagged.append({"id": r.get("id"), "statement": r.get("statement"), "unseen": sorted(unseen)})
    return flagged


def check(b, spec, src_root):
    ans_by_book = answers(spec["home"])
    by_q = {}
    for book in ans_by_book.values():
        by_q.update(book)
    res = {"business": b, "files": [], "substantive": {k: interview.substantive(v) for k, v in ans_by_book.items()},
           "prompts": {k: len(v) for k, v in ans_by_book.items()}}
    told_total = 0
    for p in spec["brain_files"]:
        f = {"file": os.path.basename(p), "exists": os.path.exists(p)}
        if f["exists"]:
            recs, _w, info = brainzip.read_brain(p)
            told = [r for r in recs if r.get("source") == "told"]
            told_total += len(told)
            f.update(brain_tab=bool(info.get("present")), records=len(recs), told=len(told),
                     lint=lint(recs, by_q))
            ok, why = same_data(os.path.join(src_root, b, os.path.basename(p)), p)
            f.update(data_intact=ok, data_why=why)
        res["files"].append(f)
    res["pass_code_gate"] = (told_total >= 1 and all(f.get("exists") and f.get("brain_tab") and f.get("data_intact")
                                                     for f in res["files"]))
    res["lint_flags"] = sum(len(f.get("lint") or []) for f in res["files"])
    res["max_substantive"] = max(res["substantive"].values() or [0])
    res["max_prompts"] = max(res["prompts"].values() or [0])
    return res


if __name__ == "__main__":
    spec = json.load(open(sys.argv[1]))
    out = [check(b, spec[b], sys.argv[2]) for b in spec]
    print(json.dumps(out, indent=1, ensure_ascii=False, default=str))

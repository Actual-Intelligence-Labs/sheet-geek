"""Load a workbook or CSV into plain Python rows, with formulas kept separately.

Code reads rows; the model only ever sees profiles. Runtime needs openpyxl for
.xlsx/.xlsm and nothing but the stdlib for .csv.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import os
import re
from dataclasses import dataclass, field

from . import brainzip

MAX_ROWS = 2_000_000


@dataclass
class Sheet:
    name: str
    state: str = "visible"
    values: list = field(default_factory=list)      # rows of cached values
    formulas: dict = field(default_factory=dict)    # (r, c) -> "=..."
    merged: list = field(default_factory=list)      # (r1, c1, r2, c2), 0-based inclusive
    missing_cached: int = 0
    truncated: bool = False
    is_brain: bool = False
    is_rules: bool = False

    @property
    def n_rows(self) -> int:
        return len(self.values)

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.values), default=0)


@dataclass
class Book:
    path: str
    kind: str                      # xlsx | csv
    sheets: list = field(default_factory=list)
    pkg: object = None             # brainzip.Package for xlsx
    csv_dialect: str = ""
    encoding: str = ""
    size: int = 0

    @property
    def name(self) -> str:
        return os.path.basename(self.path)

    def data_sheets(self) -> list:
        return [s for s in self.sheets if not s.is_brain and not s.is_rules]


def load(path: str) -> Book:
    path = os.path.abspath(path)
    low = path.lower()
    if low.endswith((".csv", ".tsv", ".txt")):
        return _load_csv(path)
    if low.endswith((".xlsx", ".xlsm")):
        return _load_xlsx(path)
    if low.endswith(".xls"):
        raise brainzip.BrainError("Old .xls files are not supported. Save it as .xlsx first.")
    raise brainzip.BrainError(f"Unsupported file type: {os.path.basename(path)}")


def _trim(row: list) -> list:
    i = len(row)
    while i and (row[i - 1] is None or row[i - 1] == ""):
        i -= 1
    return list(row[:i])


def _load_xlsx(path: str) -> Book:
    import warnings
    try:
        import openpyxl  # noqa: PLC0415  (only needed for workbooks)
    except ImportError:
        raise brainzip.BrainError("Reading .xlsx files needs the openpyxl package. Install it with: "
                                  "python3 -m pip install openpyxl") from None

    warnings.filterwarnings("ignore", module="openpyxl")

    pkg = brainzip.load_package(path)
    book = Book(path, "xlsx", pkg=pkg, size=os.path.getsize(path))
    states = {s.name: s.state for s in pkg.sheets}
    parts = {s.name: s.part for s in pkg.sheets}
    wb_v = openpyxl.load_workbook(path, read_only=True, data_only=True)
    wb_f = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        for ws in wb_v.worksheets:
            sh = Sheet(ws.title, states.get(ws.title, "visible"))
            sh.is_brain = ws.title.lower() == brainzip.BRAIN_SHEET
            sh.is_rules = ws.title.strip().lower() == ".rules"
            if sh.is_brain:
                book.sheets.append(sh)
                continue
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS:
                    sh.truncated = True
                    break
                rows.append(_trim(list(row)))
            while rows and not rows[-1]:
                rows.pop()
            sh.values = rows
            wsf = wb_f[ws.title]
            for r, row in enumerate(wsf.iter_rows(values_only=True)):
                if r >= len(rows):
                    break
                for c, v in enumerate(row):
                    f = _formula_text(v)
                    if f is not None:
                        sh.formulas[(r, c)] = f
                        if c >= len(rows[r]) or rows[r][c] is None:
                            sh.missing_cached += 1
            part = parts.get(ws.title)
            if part and part in pkg.data:
                sh.merged = _merged_ranges(pkg.data[part])
            book.sheets.append(sh)
    finally:
        wb_v.close()
        wb_f.close()
    return book


def _formula_text(v) -> str | None:
    if isinstance(v, str):
        return v if v.startswith("=") and len(v) > 1 else None
    text = getattr(v, "text", None)          # ArrayFormula
    if isinstance(text, str) and text.startswith("="):
        return text
    return None


_MERGE = re.compile(rb'<(?:\w+:)?mergeCell\s+ref="([A-Z]+)(\d+):([A-Z]+)(\d+)"')


def _merged_ranges(xml: bytes) -> list:
    out = []
    for m in _MERGE.finditer(xml):
        c1, r1, c2, r2 = m.groups()
        out.append((int(r1) - 1, brainzip.col_index(c1.decode()),
                    int(r2) - 1, brainzip.col_index(c2.decode())))
    return out


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------
_NUM = re.compile(r"^\(?-?\$?\s*-?[\d,]*\.?\d+\)?%?$")
_DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%d-%b-%Y",
                 "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%m/%d/%Y %H:%M",
                 "%b %d, %Y", "%d/%m/%Y")


def parse_scalar(s: str):
    """Best-effort typing of a CSV string: number, date, bool-ish stays text."""
    if s is None:
        return None
    t = s.strip()
    if t == "":
        return None
    if _NUM.match(t) and any(ch.isdigit() for ch in t):
        neg = t.startswith("(") and t.endswith(")")
        core = t.strip("()").replace("$", "").replace(",", "").strip()
        pct = core.endswith("%")
        core = core.rstrip("%")
        # keep leading-zero codes (00123) and long digit strings (IDs, phones) as text
        digits = core.lstrip("-")
        if (digits.startswith("0") and len(digits) > 1 and "." not in digits) or len(digits) > 15:
            return t
        try:
            n = float(core) if ("." in core or pct) else int(core)
        except ValueError:
            return t
        if pct:
            n = n / 100.0
        return -n if neg else n
    if len(t) >= 6 and any(ch.isdigit() for ch in t) and ("-" in t or "/" in t or "," in t):
        for fmt in _DATE_FORMATS:
            try:
                return dt.datetime.strptime(t, fmt)
            except ValueError:
                continue
    return t


def _load_csv(path: str) -> Book:
    raw = open(path, "rb").read()
    enc = "utf-8-sig"
    try:
        text = raw.decode(enc)
    except UnicodeDecodeError:
        enc = "latin-1"
        text = raw.decode(enc)
    sample = text[:65536]
    delim = "\t" if path.lower().endswith(".tsv") else ","
    try:
        d = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = d.delimiter
    except csv.Error:
        pass
    book = Book(path, "csv", csv_dialect=delim, encoding=enc, size=len(raw))
    sh = Sheet(os.path.splitext(os.path.basename(path))[0])
    rows = []
    for i, r in enumerate(csv.reader(io.StringIO(text), delimiter=delim)):
        if i >= MAX_ROWS:
            sh.truncated = True
            break
        typed = [parse_scalar(v) for v in r] if i else [v.strip() if isinstance(v, str) else v
                                                         for v in r]
        rows.append(_trim(typed))
    while rows and not rows[-1]:
        rows.pop()
    sh.values = rows
    book.sheets.append(sh)
    return book

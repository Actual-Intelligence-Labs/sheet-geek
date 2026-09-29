"""When a sheet with a brain comes back with more data, the brain grows: the
owner's rules cover the new rows without asking, and only what is new is asked
about, with the answer the brain already points to."""
import datetime as dt
import os
import shutil
import sys
import warnings

import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "skills", "spreadsheet-brain", "scripts"))
from sheetbrain import findings, fresh, interview, rules  # noqa: E402
from sheetbrain.analyze import Analysis  # noqa: E402
from sheetbrain.brain import Composer  # noqa: E402

FX = os.path.join(ROOT, "evals", "fixtures")


def _answer(analysis, answers, prefix, reply):
    q = next(q for q in interview.candidates(analysis, answers) if q.id.startswith(prefix))
    answers[q.id] = interview.parse_answers([q], reply)[q.id]


@pytest.fixture(scope="module")
def grown(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("grow")
    hotel, contracts = tmp / "procurement_hotel.xlsx", tmp / "procurement_contracts.xlsx"
    shutil.copy(os.path.join(FX, "procurement_hotel.xlsx"), hotel)
    shutil.copy(os.path.join(FX, "procurement_contracts.xlsx"), contracts)
    a = Analysis([str(hotel), str(contracts)])
    answers = {}
    _answer(a, answers, "find_exclusive_", "a")
    _answer(a, answers, "find_unmatched_item", "ab FEE-FUEL and FEE-DLVY are fees; the other three were re-coded")
    a.apply_answers(answers)
    recs = Composer(a, str(hotel), "b1", answers).compose()
    # a month later: more lines, a new fee code, a new hotel, more commissary transfers
    new = tmp / "grown.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(hotel)
    ws = wb["Detail"]
    hdr = [c.value for c in ws[3]]
    h = {k: i for i, k in enumerate(hdr)}
    base = [[c.value for c in r] for r in ws.iter_rows(min_row=4, max_row=40)]
    cmsy = next([c.value for c in r] for r in ws.iter_rows(min_row=4) if r[h["Location"]].value == "CMSY")
    extra = []
    for i in range(8):
        r = list(base[i]); r[h["Item #"]] = "FEE-PALLET"; r[h["Description"]] = "PALLET DEPOSIT FEE"
        r[h["Category"]] = "Fees"; extra.append(r)
    for i in range(8, 20):
        r = list(base[i]); r[h["Location"]] = "Anchor Bay"; extra.append(r)
    extra += [list(cmsy) for _ in range(5)]
    for r in extra:
        r[h["Invoice Date"]] = dt.datetime(2026, 7, 15)
    start = ws.max_row + 1
    for k, r in enumerate(extra):
        for j, v in enumerate(r):
            ws.cell(row=start + k, column=j + 1, value=v)
    wb.save(new)
    shutil.copy(new, hotel)
    a2 = Analysis([str(hotel), str(contracts)])
    report = fresh.check(a2, str(hotel), recs)
    return a2, answers, recs, report


def test_the_owners_rules_cover_the_new_rows(grown):
    a2, answers, _recs, _rep = grown
    excluded = a2.apply_answers(answers)
    assert any("CMSY" in vals for _s, _h, vals in excluded)          # new CMSY rows stay out, no question
    assert {"fee-fuel", "fee-dlvy"} <= rules.not_items(a2, answers)
    assert "fee-pallet" not in rules.not_items(a2, answers)           # new codes are asked about, not assumed


def test_only_what_is_new_is_asked_with_a_suggestion(grown):
    a2, answers, recs, report = grown
    qs = findings.grow_questions(a2, answers, report, since="2026-09-25", prev=recs)
    fee = next(q for q in qs if "FEE-PALLET" in q.prompt)
    assert fee.recommend == "not_items" and "FEE-" in fee.recommend_basis
    hotel = next(q for q in qs if "Anchor Bay" in q.prompt)
    assert hotel.recommend is None             # nothing checked says it is a real one: no suggestion
    assert not any("FEE-FUEL" in q.prompt for q in qs)                 # already answered: never asked again


def test_a_new_code_joins_the_rules_once_answered(grown):
    a2, answers, recs, report = grown
    fee = next(q for q in findings.grow_questions(a2, answers, report, prev=recs) if "FEE-PALLET" in q.prompt)
    more = dict(answers, **{fee.id: interview.parse_answers([fee], "a")[fee.id]})
    assert "fee-pallet" in rules.not_items(a2, more)


def test_formulas_that_stop_short_of_new_rows_are_caught(grown):
    a2 = grown[0]
    short = [i for i in a2.insights if i["recipe"].startswith("formula:short_range:")]
    assert short and "Summary" in short[0]["statement"] and "left out" in short[0]["statement"]


def test_the_tab_says_how_it_grows_right_under_its_first_row(grown):
    from sheetbrain.brain import for_tab, is_imperative
    _a2, _answers, recs, _report = grown
    tab = for_tab(recs)
    assert tab[0]["record"] == "meta" and "Written on" in tab[0]["statement"]
    assert "rules that also cover rows added later" in tab[0]["statement"]
    assert tab[1]["id"] == "f:howto" and "new rows with a newer date" in tab[1]["statement"]
    assert not is_imperative(tab[1]["statement"])                    # described, never commanded
    assert tab[2]["source"] == "told"                                 # the owner's notes come right after


def test_rows_another_ai_added_read_the_same_way():
    from sheetbrain.brainzip import tidy
    rows = tidy([{"record": "meta", "id": "brain:1", "statement": "first"},
                 {"record": "update", "label": "Rows", "statement": "Detail now has 6,315 rows.", "source": "told"},
                 {"record": "meta", "id": "brain:1:u1", "statement": "second meta"}])
    assert [r["record"] for r in rows] == ["meta", "fact", "fact"]
    assert rows[1]["kind"] == "update" and rows[1]["id"].startswith("x:")
    assert rows[1]["status"] == "unconfirmed"                          # no speaker named: never the owner's
    assert tidy(rows)[1]["id"] == rows[1]["id"]                        # the same words, the same id


def test_the_row_that_holds_now_wins():
    from sheetbrain.brain import current
    recs = [{"id": "f:a", "statement": "old", "status": "superseded", "as_of": "2026-09-01"},
            {"id": "f:b", "statement": "only old", "status": "superseded", "as_of": "2026-09-01"},
            {"id": "f:a", "statement": "new", "status": "current", "as_of": "2026-09-26"}]
    got = {r["id"]: r["statement"] for r in current(recs)}
    assert got == {"f:a": "new", "f:b": "only old"}


def test_old_counts_drop_out_but_the_owners_words_stay(grown):
    from sheetbrain import say
    _a2, _answers, recs, _report = grown
    told = next(r for r in recs if r.get("source") == "told" and r.get("record") == "fact")
    ins = next(r for r in recs if r.get("record") == "insight")
    later = [dict(r, status="superseded") if r in (told, ins) else r for r in recs]
    pack = say.context_pack("x.xlsx", later, {}, origin="received")
    assert told["statement"] in pack.split("marked superseded (they may still hold):")[1]
    assert ins["statement"] not in pack and "counted note marked superseded was left out" in pack
